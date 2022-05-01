# This is a sample Python script.

# Press Shift+F10 to execute it or replace it with your code.
# Press Double Shift to search everywhere for classes, files, tool windows, actions, and settings.

from docx import Document
import os
import datetime
from docx2pdf import convert
from docx.enum.dml import MSO_THEME_COLOR
from docx.shared import RGBColor


office_month_report_template = "office_month_report_template.docx"

def get_report_template_path() -> str:
    onedrive_path = os.environ["Onedrive"]
    template_path = os.path.join(onedrive_path, "Дюс", "Lichba LLC", "Daikin", "reports", "templates")
    return template_path

def get_office_month_template() -> Document:
    template_path = get_report_template_path()
    office_month_template_path = os.path.join(template_path,office_month_report_template)
    #office_month_template_path = os.path.join(template_path, "template.docx")
    doc = Document(office_month_template_path)
    return doc

def get_report_name() -> str:
    return os.path.join(get_report_template_path(), "template1.docx")

def get_table(doc: Document, index):
    return doc.tables[index]


# Press the green button in the gutter to run the script.
if __name__ == '__main__':
    doc = get_office_month_template()
    table_HVAC = get_table(doc, 0)
    table_HVAC.add_row()
    table_HVAC.cell(2, 0).text = datetime.datetime.now().strftime("%m/%d/%yy")

    #table_HVAC.cell(2, 0).style = table_HVAC.cell(1, 0).style
    hr_table = get_table(doc, 1)
    hr_table.add_row()
    productivity_table = get_table(doc, 2)
    productivity_table.add_row()
    name = get_report_name()
    doc.save(name)
    pdf_name = name.replace('docx', 'pdf')
    #convert(name, pdf_name)

    print("doc is ready")


# See PyCharm help at https://www.jetbrains.com/help/pycharm/
